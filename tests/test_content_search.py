from pathlib import Path
import shutil
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from openpyxl import Workbook
from reportlab.pdfgen import canvas

from file_compressor.content_search import (
    SearchResult,
    _extract_hwp5_text_records,
    extract_text_from_file,
    kind_for_path,
    run_tesseract_ocr,
    search_file,
    search_files,
)
from file_compressor.dependencies import TESSERACT_DOWNLOAD_URL, detect_tesseract


def case_dir(name: str) -> Path:
    path = Path(".worktrees/file-compressor-impl/.test-output") / name
    if path.exists():
        shutil.rmtree(path)
    path.mkdir(parents=True)
    return path


def write_text_pdf(path: Path, text: str):
    pdf = canvas.Canvas(str(path))
    pdf.drawString(72, 720, text)
    pdf.save()


def write_hwpx(path: Path, text: str):
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        archive.writestr(
            "Contents/section0.xml",
            (
                '<?xml version="1.0" encoding="UTF-8"?>'
                '<hp:sec xmlns:hp="http://www.hancom.co.kr/hwpml/2011/paragraph">'
                f"<hp:p><hp:run><hp:t>{text}</hp:t></hp:run></hp:p>"
                "</hp:sec>"
            ),
        )


def write_pptx(path: Path, slide_texts: list[str]):
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for index, text in enumerate(slide_texts, start=1):
            archive.writestr(
                f"ppt/slides/slide{index}.xml",
                (
                    '<?xml version="1.0" encoding="UTF-8"?>'
                    '<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
                    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
                    f"<p:cSld><p:spTree><p:sp><p:txBody><a:p><a:r><a:t>{text}</a:t></a:r></a:p>"
                    "</p:txBody></p:sp></p:spTree></p:cSld></p:sld>"
                ),
            )


def hwp_text_record(text: str) -> bytes:
    payload = text.encode("utf-16le")
    header = 67 | (len(payload) << 20)
    return header.to_bytes(4, "little") + payload


def test_extract_hwp5_text_records_decodes_paragraph_text():
    text = _extract_hwp5_text_records(hwp_text_record("alpha needle 한글"))

    assert "alpha needle 한글" in text


def test_extract_text_from_supported_file_types():
    workdir = case_dir("content-extract")
    text_file = workdir / "memo.txt"
    pdf_file = workdir / "report.pdf"
    docx_file = workdir / "letter.docx"
    xlsx_file = workdir / "sheet.xlsx"
    hwpx_file = workdir / "draft.hwpx"
    pptx_file = workdir / "deck.pptx"
    text_file.write_text("alpha text", encoding="utf-8")
    write_text_pdf(pdf_file, "bravo pdf text")
    document = Document()
    document.add_paragraph("charlie docx text")
    document.save(docx_file)
    workbook = Workbook()
    workbook.active["A1"] = "delta xlsx text"
    workbook.save(xlsx_file)
    write_hwpx(hwpx_file, "echo hwpx text")
    write_pptx(pptx_file, ["foxtrot first slide", "golf second slide"])

    assert "alpha text" in "\n".join(section.text for section in extract_text_from_file(text_file))
    assert "bravo pdf text" in "\n".join(section.text for section in extract_text_from_file(pdf_file))
    assert "charlie docx text" in "\n".join(section.text for section in extract_text_from_file(docx_file))
    assert "delta xlsx text" in "\n".join(section.text for section in extract_text_from_file(xlsx_file))
    assert "echo hwpx text" in "\n".join(section.text for section in extract_text_from_file(hwpx_file))
    pptx_sections = extract_text_from_file(pptx_file)
    assert [section.kind for section in pptx_sections] == ["PPTX", "PPTX"]
    assert "foxtrot first slide" in "\n".join(section.text for section in pptx_sections)
    assert "golf second slide" in "\n".join(section.text for section in pptx_sections)


def test_search_files_finds_text_inside_pptx_slides():
    workdir = case_dir("content-pptx-search")
    pptx_file = workdir / "deck.pptx"
    write_pptx(pptx_file, ["intro slide", "다기능 appears in slide two"])

    results = search_files([pptx_file], "다기능")

    assert kind_for_path(pptx_file) == "PPTX"
    assert results == [
        SearchResult(
            source=pptx_file,
            kind="PPTX",
            location="Slide 2",
            snippet="다기능 appears in slide two",
            status="matched",
        )
    ]


def test_search_files_uses_hancom_automation_for_hwp_content():
    workdir = case_dir("content-hwp")
    source = workdir / "sample.hwp"
    source.write_bytes(b"hwp")
    events = []

    class FakeHwp:
        def RegisterModule(self, dll_name, module_name):
            events.append(("register", dll_name, module_name))

        def Open(self, target):
            events.append(("open", Path(target)))
            return True

        def GetTextFile(self, format_name, options):
            events.append(("get_text_file", format_name, options))
            return "alpha needle from hwp"

        def Quit(self):
            events.append(("quit",))

    results = search_files(
        [source],
        "needle",
        hwp_automation_available=lambda: True,
        hwp_dispatch=lambda progid: FakeHwp(),
    )

    assert results == [
        SearchResult(
            source=source,
            kind="HWP",
            location="Document",
            snippet="alpha needle from hwp",
            status="matched",
        )
    ]
    assert events[0][0] == "register"
    assert any(event[0] == "open" for event in events)
    assert ("get_text_file", "TEXT", "") in events
    assert ("quit",) in events


def test_hwp_open_retries_with_format_arguments_when_single_path_open_fails():
    workdir = case_dir("content-hwp-open-format")
    source = workdir / "sample.hwp"
    source.write_bytes(b"hwp")
    events = []

    class FakeHwp:
        def RegisterModule(self, dll_name, module_name):
            pass

        def Open(self, *args):
            events.append(("open", args))
            if len(args) == 1:
                raise RuntimeError("wrong number of parameters")
            return True

        def GetTextFile(self, format_name, options):
            return "needle after formatted open"

        def Quit(self):
            pass

    results = search_file(
        source,
        "needle",
        hwp_automation_available=lambda: True,
        hwp_dispatch=lambda progid: FakeHwp(),
    )

    assert results[0].status == "matched"
    assert events == [
        ("open", (str(source.resolve()),)),
        ("open", (str(source.resolve()), "HWP", "")),
    ]


def test_search_file_accepts_hwp_get_text_file_with_one_argument():
    workdir = case_dir("content-hwp-gettextfile-one-arg")
    source = workdir / "sample.hwp"
    source.write_bytes(b"hwp")
    calls = []

    class FakeHwp:
        def RegisterModule(self, dll_name, module_name):
            pass

        def Open(self, target):
            return True

        def GetTextFile(self, format_name):
            calls.append(format_name)
            return "needle from one argument"

        def Quit(self):
            pass

    results = search_file(
        source,
        "needle",
        hwp_automation_available=lambda: True,
        hwp_dispatch=lambda progid: FakeHwp(),
    )

    assert results[0].status == "matched"
    assert results[0].snippet == "needle from one argument"
    assert calls == ["TEXT"]


def test_search_files_reports_hwp_dependency_when_hancom_is_unavailable():
    workdir = case_dir("content-hwp-missing")
    source = workdir / "sample.hwp"
    source.write_bytes(b"hwp")

    results = search_files([source], "needle", hwp_automation_available=lambda: False)

    assert results[0].kind == "HWP"
    assert results[0].status == "skipped"
    assert "Hancom Office" in results[0].message


def test_search_files_reports_matches_with_locations_and_snippets():
    workdir = case_dir("content-search")
    source = workdir / "memo.txt"
    source.write_text("first line\nneedle appears here\nlast line", encoding="utf-8")

    results = search_files([source], "needle")

    assert results == [
        SearchResult(
            source=source,
            kind="Text",
            location="Line 2",
            snippet="needle appears here",
            status="matched",
        )
    ]


def test_search_files_reports_no_match_and_unsupported_files():
    workdir = case_dir("content-no-match")
    text_file = workdir / "memo.txt"
    binary_file = workdir / "app.bin"
    text_file.write_text("nothing useful", encoding="utf-8")
    binary_file.write_bytes(b"\x00\x01")

    results = search_files([text_file, binary_file], "needle")

    assert results[0].status == "no_match"
    assert results[1].status == "skipped"
    assert results[1].message == "Unsupported file type."


def test_detect_tesseract_and_download_url():
    assert detect_tesseract(lambda name: "C:/Tools/tesseract.exe" if name == "tesseract" else None).available is True
    assert detect_tesseract(lambda name: None).available is False
    assert TESSERACT_DOWNLOAD_URL == "https://github.com/UB-Mannheim/tesseract/wiki"


def test_run_tesseract_ocr_uses_runner_and_search_can_use_ocr_fallback():
    workdir = case_dir("content-ocr")
    image = workdir / "scan.png"
    image.write_bytes(b"fake image")
    calls = []

    def fake_runner(command, **kwargs):
        calls.append((command, kwargs))

        class Result:
            stdout = "needle from ocr"
            stderr = ""
            returncode = 0

        return Result()

    assert run_tesseract_ocr(image, executable="tesseract", runner=fake_runner) == "needle from ocr"

    results = search_files(
        [image],
        "needle",
        use_ocr=True,
        tesseract_executable="tesseract",
        ocr_runner=fake_runner,
    )

    assert calls[0][0] == ["tesseract", str(image), "stdout", "-l", "kor+eng"]
    assert results[0].status == "matched"
    assert results[0].location == "OCR"


def test_search_files_passes_selected_ocr_language_to_tesseract():
    workdir = case_dir("content-ocr-language")
    image = workdir / "scan.png"
    image.write_bytes(b"fake image")
    calls = []

    def fake_runner(command, **kwargs):
        calls.append(command)

        class Result:
            stdout = "needle from ocr"
            stderr = ""
            returncode = 0

        return Result()

    results = search_files(
        [image],
        "needle",
        use_ocr=True,
        tesseract_executable="tesseract",
        ocr_language="eng",
        ocr_runner=fake_runner,
    )

    assert calls == [["tesseract", str(image), "stdout", "-l", "eng"]]
    assert results[0].status == "matched"


def test_pdf_ocr_renders_pages_before_running_tesseract():
    workdir = case_dir("content-pdf-ocr-render")
    pdf = workdir / "scan.pdf"
    write_text_pdf(pdf, "no searchable OCR keyword")
    calls = []
    rendered_paths = []

    def fake_renderer(source, output_dir):
        assert source == pdf
        first = output_dir / "page-0001.png"
        second = output_dir / "page-0002.png"
        first.write_bytes(b"first")
        second.write_bytes(b"second")
        rendered_paths.extend([first, second])
        return [first, second]

    def fake_runner(command, **kwargs):
        calls.append(command)

        class Result:
            stderr = ""
            returncode = 0

            @property
            def stdout(self):
                return "needle on rendered page" if "page-0002.png" in command[1] else "no match"

        return Result()

    results = search_file(
        pdf,
        "needle",
        use_ocr=True,
        tesseract_executable="tesseract",
        ocr_language="eng",
        ocr_runner=fake_runner,
        pdf_page_renderer=fake_renderer,
    )

    assert calls == [
        ["tesseract", str(rendered_paths[0]), "stdout", "-l", "eng"],
        ["tesseract", str(rendered_paths[1]), "stdout", "-l", "eng"],
    ]
    assert results == [
        SearchResult(
            source=pdf,
            kind="OCR",
            location="OCR Page 2",
            snippet="needle on rendered page",
            status="matched",
        )
    ]
