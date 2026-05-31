from __future__ import annotations

from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path
import subprocess
from typing import Any
import xml.etree.ElementTree as ET
from zipfile import ZipFile
import zlib

from docx import Document
import olefile
from openpyxl import load_workbook
from pypdf import PdfReader

from file_compressor.compressors.windows_automation import (
    HWP_APPLICATION_PROGID,
    _initialize_com,
    default_com_dispatch,
    default_hancom_available,
)


TEXT_SUFFIXES = {".txt", ".csv", ".md", ".log"}
PDF_SUFFIXES = {".pdf"}
DOCX_SUFFIXES = {".docx"}
XLSX_SUFFIXES = {".xlsx", ".xlsm"}
HWPX_SUFFIXES = {".hwpx"}
HWP_SUFFIXES = {".hwp"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
HWP5_PARA_TEXT_TAG = 67


class ContentSearchDependencyError(RuntimeError):
    pass


@dataclass(frozen=True)
class TextSection:
    kind: str
    location: str
    text: str


@dataclass(frozen=True)
class SearchResult:
    source: Path
    kind: str
    location: str
    snippet: str
    status: str
    message: str = ""


def extract_text_from_file(
    path: Path,
    *,
    hwp_automation_available: Callable[[], bool] = default_hancom_available,
    hwp_dispatch: Callable[[str], Any] = default_com_dispatch,
) -> list[TextSection]:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in TEXT_SUFFIXES:
        return _extract_text_file(path)
    if suffix in PDF_SUFFIXES:
        return _extract_pdf(path)
    if suffix in DOCX_SUFFIXES:
        return _extract_docx(path)
    if suffix in XLSX_SUFFIXES:
        return _extract_xlsx(path)
    if suffix in HWPX_SUFFIXES:
        return _extract_hwpx(path)
    if suffix in HWP_SUFFIXES:
        return _extract_hwp(
            path,
            automation_available=hwp_automation_available,
            dispatch=hwp_dispatch,
        )
    raise ValueError("Unsupported file type.")


def search_files(
    paths: list[Path],
    query: str,
    *,
    use_ocr: bool = False,
    tesseract_executable: str | None = None,
    ocr_runner=subprocess.run,
    hwp_automation_available: Callable[[], bool] = default_hancom_available,
    hwp_dispatch: Callable[[str], Any] = default_com_dispatch,
) -> list[SearchResult]:
    results: list[SearchResult] = []
    for path in paths:
        results.extend(
            search_file(
                path,
                query,
                use_ocr=use_ocr,
                tesseract_executable=tesseract_executable,
                ocr_runner=ocr_runner,
                hwp_automation_available=hwp_automation_available,
                hwp_dispatch=hwp_dispatch,
            )
        )
    return results


def search_file(
    path: Path,
    query: str,
    *,
    use_ocr: bool = False,
    tesseract_executable: str | None = None,
    ocr_runner=subprocess.run,
    hwp_automation_available: Callable[[], bool] = default_hancom_available,
    hwp_dispatch: Callable[[str], Any] = default_com_dispatch,
) -> list[SearchResult]:
    path = Path(path)
    query_normalized = query.lower()
    try:
        sections = extract_text_from_file(
            path,
            hwp_automation_available=hwp_automation_available,
            hwp_dispatch=hwp_dispatch,
        )
    except ValueError:
        if use_ocr and path.suffix.lower() in IMAGE_SUFFIXES | PDF_SUFFIXES:
            return _search_ocr(path, query_normalized, tesseract_executable, ocr_runner)
        return [
            SearchResult(
                source=path,
                kind=kind_for_path(path),
                location="-",
                snippet="",
                status="skipped",
                message="Unsupported file type.",
            )
        ]
    except ContentSearchDependencyError as exc:
        return [
            SearchResult(
                source=path,
                kind=kind_for_path(path),
                location="-",
                snippet="",
                status="skipped",
                message=str(exc),
            )
        ]
    except Exception as exc:
        message = _exception_message(exc)
        if path.suffix.lower() in HWP_SUFFIXES:
            message = f"HWP text extraction failed: {message}"
        return [
            SearchResult(
                source=path,
                kind=kind_for_path(path),
                location="-",
                snippet="",
                status="failed",
                message=message,
            )
        ]

    matched = _search_sections(path, sections, query_normalized)
    if matched:
        return matched
    if use_ocr and path.suffix.lower() in IMAGE_SUFFIXES | PDF_SUFFIXES:
        return _search_ocr(path, query_normalized, tesseract_executable, ocr_runner)
    return [
        SearchResult(
            source=path,
            kind=kind_for_path(path),
            location="-",
            snippet="",
            status="no_match",
        )
    ]


def _exception_message(exc: Exception) -> str:
    message = str(exc).strip()
    if message:
        return message
    return exc.__class__.__name__


def run_tesseract_ocr(
    path: Path,
    *,
    executable: str = "tesseract",
    language: str = "kor+eng",
    runner=subprocess.run,
) -> str:
    command = [executable, str(path), "stdout", "-l", language]
    result = runner(command, capture_output=True, text=True, encoding="utf-8", errors="replace")
    if result.returncode != 0:
        raise OSError(result.stderr.strip() or "Tesseract OCR failed.")
    return result.stdout


def _extract_text_file(path: Path) -> list[TextSection]:
    text = path.read_text(encoding="utf-8", errors="ignore")
    sections: list[TextSection] = []
    for index, line in enumerate(text.splitlines(), start=1):
        sections.append(TextSection(kind="Text", location=f"Line {index}", text=line))
    return sections


def _extract_pdf(path: Path) -> list[TextSection]:
    reader = PdfReader(str(path))
    sections: list[TextSection] = []
    for index, page in enumerate(reader.pages, start=1):
        sections.append(TextSection(kind="PDF", location=f"Page {index}", text=page.extract_text() or ""))
    return sections


def _extract_docx(path: Path) -> list[TextSection]:
    document = Document(str(path))
    sections: list[TextSection] = []
    for index, paragraph in enumerate(document.paragraphs, start=1):
        if paragraph.text:
            sections.append(TextSection(kind="DOCX", location=f"Paragraph {index}", text=paragraph.text))
    return sections


def _extract_xlsx(path: Path) -> list[TextSection]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sections: list[TextSection] = []
    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        sections.append(
                            TextSection(
                                kind="XLSX",
                                location=f"{sheet.title}!{cell.coordinate}",
                                text=str(cell.value),
                            )
                        )
    finally:
        workbook.close()
    return sections


def _extract_hwpx(path: Path) -> list[TextSection]:
    sections: list[TextSection] = []
    with ZipFile(path, "r") as archive:
        for name in sorted(archive.namelist()):
            lower_name = name.lower()
            if not lower_name.startswith("contents/") or not lower_name.endswith(".xml"):
                continue
            text = _xml_text(archive.read(name))
            if text:
                sections.append(TextSection(kind="HWPX", location=name, text=text))
    return sections


def _xml_text(data: bytes) -> str:
    root = ET.fromstring(data)
    return " ".join(part.strip() for part in root.itertext() if part.strip())


def _extract_hwp(
    path: Path,
    *,
    automation_available: Callable[[], bool],
    dispatch: Callable[[str], Any],
) -> list[TextSection]:
    try:
        sections = _extract_hwp5(path)
    except ContentSearchDependencyError:
        raise
    except Exception:
        sections = []
    if sections:
        return sections

    if not automation_available():
        raise ContentSearchDependencyError("Hancom Office is required for HWP content search.")

    app = None
    uninitialize = _initialize_com()
    try:
        app = dispatch(HWP_APPLICATION_PROGID)
        _open_hwp_document(app, path)
        text = _read_hwp_text(app)
    finally:
        try:
            if app is not None:
                app.Quit()
        finally:
            uninitialize()

    if not text:
        return []
    return [TextSection(kind="HWP", location="Document", text=text)]


def _extract_hwp5(path: Path) -> list[TextSection]:
    if not olefile.isOleFile(str(path)):
        return []

    sections: list[TextSection] = []
    with olefile.OleFileIO(str(path)) as ole:
        if not ole.exists("FileHeader"):
            return []
        header = ole.openstream("FileHeader").read()
        properties = int.from_bytes(header[36:40], "little") if len(header) >= 40 else 0
        if properties & 0x02:
            raise ContentSearchDependencyError("Password-protected HWP files cannot be searched.")
        compressed = bool(properties & 0x01)
        section_paths = sorted(
            (
                item
                for item in ole.listdir(streams=True, storages=False)
                if len(item) == 2 and item[0] == "BodyText" and item[1].startswith("Section")
            ),
            key=lambda item: _hwp_section_index(item[1]),
        )
        for section_path in section_paths:
            data = ole.openstream(section_path).read()
            if compressed:
                data = _decompress_hwp_stream(data)
            text = _extract_hwp5_text_records(data)
            if text:
                sections.append(
                    TextSection(
                        kind="HWP",
                        location="/".join(section_path),
                        text=text,
                    )
                )
    return sections


def _hwp_section_index(name: str) -> int:
    try:
        return int(name.removeprefix("Section"))
    except ValueError:
        return 0


def _decompress_hwp_stream(data: bytes) -> bytes:
    try:
        return zlib.decompress(data, -15)
    except zlib.error:
        return zlib.decompress(data)


def _extract_hwp5_text_records(data: bytes) -> str:
    offset = 0
    paragraphs: list[str] = []
    while offset + 4 <= len(data):
        header = int.from_bytes(data[offset : offset + 4], "little")
        offset += 4
        tag_id = header & 0x3FF
        size = (header >> 20) & 0xFFF
        if size == 0xFFF:
            if offset + 4 > len(data):
                break
            size = int.from_bytes(data[offset : offset + 4], "little")
            offset += 4
        payload = data[offset : offset + size]
        offset += size
        if tag_id == HWP5_PARA_TEXT_TAG:
            text = _decode_hwp5_text_payload(payload)
            if text:
                paragraphs.append(text)
    return "\n".join(paragraphs)


def _decode_hwp5_text_payload(payload: bytes) -> str:
    text = payload.decode("utf-16le", errors="ignore")
    return "".join(
        character if character >= " " or character in "\r\n\t" else " "
        for character in text
    ).strip()


def _open_hwp_document(app: Any, path: Path) -> None:
    try:
        app.RegisterModule("FilePathCheckDLL", "FilePathCheckerModule")
    except Exception:
        pass

    target = str(path.resolve())
    errors: list[Exception] = []
    for args in [(target,), (target, "HWP", "")]:
        try:
            opened = app.Open(*args)
        except Exception as exc:
            errors.append(exc)
            continue
        if opened is not False:
            return
        errors.append(RuntimeError("HWP file could not be opened."))
    if errors:
        raise RuntimeError(f"HWP file could not be opened: {_exception_message(errors[-1])}")
    raise RuntimeError("HWP file could not be opened.")


def _read_hwp_text(app: Any) -> str:
    text = _read_hwp_text_file(app)
    if text:
        return text
    return _scan_hwp_text(app)


def _read_hwp_text_file(app: Any) -> str:
    get_text_file = getattr(app, "GetTextFile", None)
    if not callable(get_text_file):
        return ""
    for args in [("TEXT",), ("TEXT", "")]:
        try:
            text = get_text_file(*args)
        except Exception:
            continue
        if text:
            return str(text)
        return ""
    return ""


def _scan_hwp_text(app: Any) -> str:
    try:
        app.InitScan(0x07, 0x0077, 0, 0, -1, -1)
    except Exception:
        app.InitScan()

    parts: list[str] = []
    try:
        for _ in range(100000):
            status, text = _normalize_hwp_text_result(app.GetText())
            if text:
                parts.append(text)
            if status in {0, 1}:
                break
            if status in {101, 102}:
                raise RuntimeError("HWP text scan failed.")
        else:
            raise RuntimeError("HWP text scan did not finish.")
    finally:
        try:
            app.ReleaseScan()
        except Exception:
            pass
    return "".join(parts)


def _normalize_hwp_text_result(result: Any) -> tuple[int | None, str]:
    if isinstance(result, tuple):
        if len(result) >= 2:
            return _coerce_status(result[0]), "" if result[1] is None else str(result[1])
        if len(result) == 1:
            return _coerce_status(result[0]), ""
    if result is None:
        return 0, ""
    return None, str(result)


def _coerce_status(value: Any) -> int | None:
    try:
        return int(value)
    except (TypeError, ValueError):
        return None


def _search_sections(path: Path, sections: list[TextSection], query_normalized: str) -> list[SearchResult]:
    matches: list[SearchResult] = []
    for section in sections:
        if query_normalized in section.text.lower():
            matches.append(
                SearchResult(
                    source=path,
                    kind=section.kind,
                    location=section.location,
                    snippet=_snippet(section.text, query_normalized),
                    status="matched",
                )
            )
    return matches


def _search_ocr(path: Path, query_normalized: str, executable: str | None, runner) -> list[SearchResult]:
    if not executable:
        return [
            SearchResult(
                source=path,
                kind="OCR",
                location="-",
                snippet="",
                status="skipped",
                message="OCR is not installed.",
            )
        ]
    try:
        text = run_tesseract_ocr(path, executable=executable, runner=runner)
    except OSError as exc:
        return [
            SearchResult(
                source=path,
                kind="OCR",
                location="-",
                snippet="",
                status="failed",
                message=str(exc),
            )
        ]
    if query_normalized not in text.lower():
        return [
            SearchResult(
                source=path,
                kind="OCR",
                location="OCR",
                snippet="",
                status="no_match",
            )
        ]
    return [
        SearchResult(
            source=path,
            kind="OCR",
            location="OCR",
            snippet=_snippet(text, query_normalized),
            status="matched",
        )
    ]


def _snippet(text: str, query_normalized: str) -> str:
    normalized = text.lower()
    index = normalized.find(query_normalized)
    if index < 0:
        return text.strip()[:120]
    start = max(index - 40, 0)
    end = min(index + len(query_normalized) + 40, len(text))
    return " ".join(text[start:end].strip().split())


def kind_for_path(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in TEXT_SUFFIXES:
        return "Text"
    if suffix in PDF_SUFFIXES:
        return "PDF"
    if suffix in DOCX_SUFFIXES:
        return "DOCX"
    if suffix in XLSX_SUFFIXES:
        return "XLSX"
    if suffix in HWPX_SUFFIXES:
        return "HWPX"
    if suffix in HWP_SUFFIXES:
        return "HWP"
    if suffix in IMAGE_SUFFIXES:
        return "Image"
    return "Unknown"


def _kind_for_path(path: Path) -> str:
    return kind_for_path(path)
